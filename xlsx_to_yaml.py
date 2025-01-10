import openpyxl
import yaml

def just_do_it(xlsx_file):

    __properties_list__ = ('VVK.TYPE ID', 'Name')
    __classifications_list__ = ('NSIK LF', 'NSIK LT', 'NSIK LK')
    __multilines_list__ = ('examples')

    __default_stage__ = 's3'
    __default_lod__ = 300

    __wb__ = openpyxl.load_workbook(xlsx_file)
    __sheet = __wb__.active

    __res__ = dict()
    
    c_part = '0'
    onto = 'undef'
    onto_dict = dict()

    for row in range(2, __sheet.max_row + 1):
        if __sheet.cell(row=row, column=1).value is not None:
            c_part = __sheet.cell(row=row, column=1).value            
            __res__[c_part] = dict()
            __res__[c_part]['elements'] = list()
            __res__[c_part]['stages'] = {__default_stage__: {'lod' : __default_lod__}}
            onto = __sheet.cell(row=row+1, column=2).value
                    
        onto_dict = {'classifications': dict(), 'stages' : {__default_stage__: {'properties' : list(), 'lod': __default_lod__}}}
        for col in range(3, __sheet.max_column + 1):
            
            onto = __sheet.cell(row=row, column=2).value
            onto_dict['id'] = onto

            if onto is not None:
                tmp = __sheet.cell(row=row, column=col).value
                hdr = __sheet.cell(row=1, column=col).value

                if hdr in __classifications_list__:
                    # classification items
                    onto_dict['classifications'][hdr] = tmp

                elif hdr in __properties_list__:
                    # properties / attributes
                    _string = ''
                    if tmp is not None and '\n' in tmp:
                        for i in tmp.split("\n"):
                            _string = tmp.split("\n")
                    elif tmp is not None:
                        _string = tmp

                    if _string != '':
                        onto_dict['stages'][__default_stage__]['properties'].append({hdr : {'examples' : _string}})
                    else:
                        onto_dict['stages'][__default_stage__]['properties'].append(hdr)

                elif hdr in __multilines_list__:
                    if tmp is not None and '\n' in tmp:
                        if hdr not in onto_dict.keys():
                            onto_dict[hdr] = list()
                        for i in tmp.split("\n"):
                            onto_dict[hdr].append(i.strip(', '))
                    elif tmp is not None:
                        if hdr not in onto_dict.keys():
                            onto_dict[hdr] = list()
                        onto_dict[hdr].append(tmp.strip())

                else:
                    if tmp is not None and isinstance(tmp, str):
                        onto_dict[hdr] = tmp.strip()
                    
        if onto is not None:
            __res__[c_part]['elements'].append(onto_dict)

    return __res__



if __name__ == '__main__':
    # res = just_do_it('sandbox/L95_part.xlsx')
    res = just_do_it('sandbox/Moravai.xlsx')
    out_file  = 'sandbox/Moravai.yaml'

    # for k in res.keys():
    #     for itm in res[k]['elements']:
    #         print(F"[{k}]\t {itm} of type {type(itm)}")
        
    with open(out_file, mode='w', encoding="utf-8") as of:
        yaml.dump({'loin': res}, of, indent=2, allow_unicode=True, sort_keys=False)
