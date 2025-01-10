import openpyxl
import yaml

def just_do_it(xlsx_file):

    __properties_list__ = ('VVK.TYPE ID', 'Name')
    __classifications_list__ = ('NSIKcodeLF', 'NSIKcodeLT', 'NSIKcodeLK')
    __multilines_list__ = ('examples')

    __default_stage__ = 's3'
    __default_lod__ = 300

    __mangle = {'NSIKcodeLF': 'NSIK LF', 'NSIKcodeLT':'NSIK LT', 'NSIKcodeLK':'NSIK LK', 
                'Plotis':'Ifc__.Width', 'Svoris':'Ifc__.Weigth', 'Aukštis':'Ifc__.Heigth', 'Ilgis':'Ifc__.Length', 'Storis':'Ifc__.Depth',
                'Plotas':'Ifc__.Area', 'Tūris':'Ifc__.Volume', 'Spalva':'VVK.Colour',
                'Korpusas':'VVK.Korpusas', 'Medžiagiškumas': 'VVK.Material', 'Tipas': 'VVK.State', 'Pavadinimas':'Name', 'Pavadinimas ':'Name',
                'Techninė specifikacija': 'Pset_ConstructionAdministration.SpecificationSectionNumber',
                'TYPE ID': 'Type'}

    __wb__ = openpyxl.load_workbook(xlsx_file)
    __sheet = __wb__.active

    __res__ = dict()
    
    c_part = '0'
    onto = 'undef'
    onto_dict = dict()
    att_list = list()
    prop_list = list()
    properties_list = dict()

    old_onto = None

    for row in range(24, __sheet.max_row + 1):
        if __sheet.cell(row=row, column=1).value is not None:
            c_part = __sheet.cell(row=row, column=1).value
            __res__[c_part] = dict()
            __res__[c_part]['elements'] = list()
            # __res__[c_part]['stages'] = {__default_stage__: {'lod' : __default_lod__}}
            # onto = __sheet.cell(row=row+1, column=2).value
                    
        # for col in range(9, __sheet.max_column + 1):
            
        tmp = __sheet.cell(row=row, column=2).value or None
        if tmp is not None:
            print(F"ADDING {onto_dict} to [{c_part}]['elements']")
            __res__[c_part]['elements'].append(onto_dict)
                
            # reset dictionary
            onto_dict = {'classifications': dict(), 'stages' : {__default_stage__: {'properties' : list(), 'lod': __default_lod__}}}
                
            onto = tmp

        if tmp is not None:
            
            onto_dict['id'] = __sheet.cell(row=row, column=2).value.strip()

            if __sheet.cell(row=row, column=3).value is not None:
                onto_dict['description'] = __sheet.cell(row=row, column=3).value

            if __sheet.cell(row=row, column=7).value is not None:
                onto_dict['stages'][__default_stage__]['lod'] = __sheet.cell(row=row, column=7).value
            
            if __sheet.cell(row=row, column=8).value is not None:
                onto_dict['memo'] = __sheet.cell(row=row, column=8).value

            _examples_ = __sheet.cell(row=row, column=4).value
            # onto_dict['examples'] = __sheet.cell(row=row, column=4).value

            if isinstance(_examples_, str):
                if _examples_ is not None :
                    if  '\n' in _examples_:
                    # onto_dict['stages'][__default_stage__]['properties'] = {hdr : {'examples' : list()}}
                        _str_ = list()
                        for i in _examples_.split("\n"):
                            _str_.append(i.strip(', '))
                        onto_dict['examples'] = _str_
                    elif  ', ' in _examples_:
                    # onto_dict['stages'][__default_stage__]['properties'] = {hdr : {'examples' : list()}}
                        _str_ = list()
                        for i in _examples_.split(", "):
                            _str_.append(i.strip())
                        onto_dict['examples'] = _str_
                else:
                     onto_dict['examples'] = [_examples_.strip()]

            elif _examples_ is not None:
                        # prop_list['examples'] = [_examples]
                onto_dict['examples'] = [_examples]


            _name = _descr = _units = _examples = _memo = None
            att_list = list()
            prop_list = list()

        
        else:
            _name = __sheet.cell(row=row, column=9).value
            _descr = __sheet.cell(row=row, column=10).value
            _units = __sheet.cell(row=row, column=11).value
            _examples = __sheet.cell(row=row, column=12).value
            _memo = __sheet.cell(row=row, column=13).value

            if _name in __classifications_list__:
                hdr = __mangle[_name] if _name in __mangle.keys() else _name
                # classification items
                onto_dict['classifications'][hdr] = _examples

            else:
                    
                    # properties / attributes
                    _string = ''
                    # onto_dict = {'classifications': dict(), 'stages' : {__default_stage__: {'properties' : list(), 'lod': __default_lod__}}}
                    print(F"\t[{c_part}] [{onto}]: ({_name}) {prop_list}")

                    # onto_dict['stages'][__default_stage__]['properties'] = {hdr : dict()}

                    if isinstance(_examples, str):
                        if _examples is not None and '\n' in _examples:
                            # onto_dict['stages'][__default_stage__]['properties'] = {hdr : {'examples' : list()}}
                            
                            _str_ = list()
                            print(F"\t example LIST [{c_part}] [{onto}]: ({_examples}) {prop_list}")
                            for i in _examples.split("\n"):
                            # _string = _examples.split("\n")
                                # onto_dict['stages'][__default_stage__]['properties'][hdr]['examples'].append(i.strip(', '))

                                _str_.append(i.strip(', '))
                                # prop_list['examples'].append(i.strip(', '))

                            prop_list.append({'examples' : _str_})
                            
                        elif _examples is not None:
                            print(F"\t example ANY [{c_part}] [{onto}]: ({_examples}) {prop_list}")
                            prop_list.append({'examples' : [_examples.strip()]})
                            # onto_dict['stages'][__default_stage__]['properties'] = {hdr : {'examples' : [_examples.strip()]}}
                        # onto_dict['stages'][__default_stage__]['properties'][hdr]['examples'].append(_examples.strip())
                    elif _examples is not None:
                        # prop_list['examples'] = [_examples]
                        prop_list.append({'examples' : [_examples]})
                                            
                    if _memo is not None:
                        # print(F"\t[{c_part}] [{onto}]: {hdr} ({_examples}) {prop_list}")
                        prop_list.append({'memo' : _memo})
                        
                    if _units is not None:
                        # print(F"\t[{c_part}] [{onto}]: {hdr} ({_examples}) {prop_list}")
                        prop_list.append({'units' : _units})
                        
                    if _descr is not None:
                        # print(F"\t[{c_part}] [{onto}]: {hdr} ({_examples}) {prop_list}")
                        prop_list.append({'descr' : _descr})
                    
                    hdr = __mangle[_name] if _name in __mangle.keys() else _name
                    print(F"\t[{c_part}] [{onto}] '{hdr}': {prop_list}")
                    
                    if hdr is not None and len(onto_dict)>0:
                        if hdr not in onto_dict['stages'][__default_stage__]['properties']:
                            # onto_dict['stages'][__default_stage__]['properties'].append({hdr : prop_list})
                            
                            # check if we have duplicates in properties_list
                            if hdr in properties_list.keys():
                                print(prop_list)
                                # if prop_list[hdr]['descr'] not in properties_list[hdr]['descr']:
                                #         properties_list[hdr]['descr'].append(prop_list[hdr]['descr'])
                            else:
                                properties_list[hdr] = prop_list
                            
                            # onto_dict['stages'][__default_stage__]['properties'].append(hdr)
                            onto_dict['stages'][__default_stage__]['properties'].append({hdr : prop_list})

                            
                        # else:
                        #     onto_dict['stages'][__default_stage__]['properties'][hdr] = prop_list
                        # print(F"\t ONTO ADDED [{c_part}] [{onto}]: {onto_dict['stages'][__default_stage__]['properties'][hdr]}")
        prop_list = list()
                        
    __wb__.close()
    # __res__['properties'] = properties_list
    return {'loin' : __res__, 'properties' : properties_list}



if __name__ == '__main__':
    # res = just_do_it('sandbox/L95_part.xlsx')
    # res = just_do_it('sandbox/Moravai.xlsx')
    res = just_do_it('sandbox/L95_BEP.xlsx')
    out_file  = 'sandbox/Moravai_v2.yaml'

    # for k in res.keys():
    #     for itm in res[k]['elements']:
    #         print(F"[{k}]\t {itm} of type {type(itm)}")
        
    with open(out_file, mode='w', encoding="utf-8") as of:
        yaml.dump(res, of, indent=2, allow_unicode=True, sort_keys=False)
